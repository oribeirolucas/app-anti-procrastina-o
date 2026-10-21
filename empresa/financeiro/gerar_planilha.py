from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

wb=Workbook()
H=Font(bold=True,color="FFFFFF"); HF=PatternFill("solid",fgColor="1F4E78")
IN=PatternFill("solid",fgColor="FFF2CC"); TOT=PatternFill("solid",fgColor="DDEBF7")
B=Font(bold=True); BRL='"R$" #,##0;[Red]-"R$" #,##0'
def hdr(ws,r,vals):
    for i,v in enumerate(vals,1):
        c=ws.cell(r,i,v);c.font=H;c.fill=HF;c.alignment=Alignment(horizontal="center",wrap_text=True)

# ---------- Premissas
p=wb.active;p.title="Premissas"
hdr(p,1,["Premissa","Valor","Observação"])
prem=[
("Câmbio US$ → R$",5.5,"Ajuste conforme a cotação"),
("Ticket médio mensal por cliente (R$)",470,"Mix Essencial 297 / Pro 597 / Rede 897"),
("Setup por cliente novo (R$)",497,"Metade dos clientes paga (anual isenta)"),
("% de clientes que pagam setup",0.5,""),
("Custo variável por cliente/mês (R$)",110,"WhatsApp API + IA da recepcionista + infra"),
("Taxa do gateway Asaas (% da receita)",0.035,"Média Pix/boleto/cartão"),
("Provisão IR no CPF (% do lucro)",0.275,"Carnê-leão, alíquota máxima; vira Simples com CNPJ"),
("Custo de voz por minuto (US$)",0.12,"ElevenLabs Agents + telefonia"),
("Minutos por ligação",3,""),
]
for i,(a,b,c) in enumerate(prem,2):
    p.cell(i,1,a);x=p.cell(i,2,b);x.fill=IN;p.cell(i,3,c)
    if isinstance(b,float) and b<1: x.number_format="0.0%" if b>0.02 or "%" in a else "0.000"
p["B9"].number_format="0.00"
for col,w in zip("ABC",(42,14,52)): p.column_dimensions[col].width=w
P={k:f"Premissas!$B${i}" for i,k in enumerate(["cambio","ticket","setup","pctsetup","varcli","gateway","ir","vozmin","minlig"],2)}

# ---------- Custos mensais
c=wb.create_sheet("Custos mensais")
meses=[f"Mês {i}" for i in range(1,7)]
hdr(c,1,["Categoria","Item","Unid."]+meses+["Total 6 meses","Observação"])
# rows: (cat, item, unit, values list or formula-type, obs)
fixos=[
("Ferramentas","n8n (plano Starter)","R$",[150]*6,"Já existe; ajuste ao seu plano real"),
("Ferramentas","Claude para os agentes internos","US$",[40,50,60,70,80,90],"Via créditos do n8n ou API; cresce com o volume"),
("Ferramentas","ElevenLabs (voz, plano Creator → Pro)","US$",[22,22,99,99,99,99],"Pro no mês 3 para mais minutos/vozes"),
("Ferramentas","Domínio + landing (Vercel grátis)","R$",[40,0,0,0,0,0],"Domínio anual"),
("Telefonia","Chip WhatsApp da empresa","R$",[20]*6,""),
("Telefonia","Número BR para ligações (Twilio/SIP)","R$",[0,50,50,50,50,50],"A partir da Fase 2"),
("Telefonia","Ligações de venda (nº de ligações)","lig.",[0,100,300,500,600,600],"Edite o volume; custo calculado abaixo"),
("Mídia","Meta Ads","R$",[1500,2500,4000,6000,8000,9000],"Principal alavanca; só sobe com CAC na meta"),
("Administrativo","Contador (após CNPJ)","R$",[0,0,300,300,300,300],"CNPJ previsto no fim do mês 2"),
("Administrativo","Abertura do CNPJ (único)","R$",[0,0,400,0,0,0],"Estimativa"),
("Reserva","Imprevistos (10% dos custos fixos acima)","R$",None,"Calculado"),
]
r=2;rows={}
for cat,item,un,vals,obs in fixos:
    c.cell(r,1,cat);c.cell(r,2,item);c.cell(r,3,un);c.cell(r,11,obs)
    rows[item]=r
    if vals:
        for j,v in enumerate(vals):
            x=c.cell(r,4+j,v);x.fill=IN;x.number_format=BRL if un=="R$" else "#,##0"
    r+=1
# convert USD & ligações into R$ rows
conv_start=r
c.cell(r,1,"Conversões");c.cell(r,2,"Claude (R$)");c.cell(r,3,"R$")
for j in range(6):
    col=L(4+j);c.cell(r,4+j,f"={col}{rows['Claude para os agentes internos']}*{P['cambio']}")
r+=1
c.cell(r,1,"Conversões");c.cell(r,2,"ElevenLabs (R$)");c.cell(r,3,"R$")
for j in range(6):
    col=L(4+j);c.cell(r,4+j,f"={col}{rows['ElevenLabs (voz, plano Creator → Pro)']}*{P['cambio']}")
r+=1
c.cell(r,1,"Conversões");c.cell(r,2,"Custo das ligações (R$)");c.cell(r,3,"R$")
for j in range(6):
    col=L(4+j);c.cell(r,4+j,f"={col}{rows['Ligações de venda (nº de ligações)']}*{P['minlig']}*{P['vozmin']}*{P['cambio']}")
r+=1
conv_end=r-1
# reserva
rr=rows["Imprevistos (10% dos custos fixos acima)"]
brl_rows=[rows[k] for k,(_,_,u,_,_) in zip(rows,fixos) if u=="R$" and k!="Imprevistos (10% dos custos fixos acima)"]
for j in range(6):
    col=L(4+j)
    terms="+".join(f"{col}{x}" for x in brl_rows if x!=rows["Meta Ads"])
    c.cell(rr,4+j,f"=0.1*({terms}+SUM({col}{conv_start}:{col}{conv_end}))")
r+=1
# totals
def totrow(label,formula_fn,bold=True):
    global r
    c.cell(r,2,label).font=B
    for j in range(6):
        col=L(4+j);x=c.cell(r,4+j,"="+formula_fn(col));x.font=B;x.fill=TOT;x.number_format=BRL
    c.cell(r,10,f"=SUM(D{r}:I{r})").font=B;c.cell(r,10).number_format=BRL;c.cell(r,10).fill=TOT
    r+=1;return r-1
brl_all=brl_rows+[rr]
tot_op=totrow("CUSTO OPERACIONAL (sem mídia)",lambda col:"+".join(f"{col}{x}" for x in brl_all if x!=rows["Meta Ads"])+f"+SUM({col}{conv_start}:{col}{conv_end})")
tot_mid=totrow("MÍDIA",lambda col:f"{col}{rows['Meta Ads']}")
tot_all=totrow("CUSTO FIXO TOTAL DO MÊS",lambda col:f"{col}{tot_op}+{col}{tot_mid}")
for rr_ in range(2,r):
    for j in range(6): 
        cell=c.cell(rr_,4+j)
        if cell.number_format=="General": cell.number_format=BRL
    if c.cell(rr_,10).value is None and rr_<tot_op: 
        c.cell(rr_,10,f'=IF(C{rr_}="R$",SUM(D{rr_}:I{rr_}),"")').number_format=BRL
for col,w in zip("ABCDEFGHIJK",(15,40,7,11,11,11,11,11,11,14,48)): c.column_dimensions[col].width=w
c.freeze_panes="D2"

# ---------- Resultado e caixa
f=wb.create_sheet("Resultado e caixa")
hdr(f,1,["Linha"]+meses+["Total"])
CM="'Custos mensais'!"
lines=[
("Clientes ativos (fim do mês)",[8,22,40,62,90,107],"in"),
("Clientes novos",None,"novos"),
("Receita recorrente (MRR)",None,"mrr"),
("Receita de setup",None,"setup"),
("RECEITA TOTAL",None,"rec"),
("(-) Custo variável dos clientes",None,"var"),
("(-) Taxa do gateway",None,"gw"),
("(-) Custo operacional",None,"op"),
("(-) Mídia",None,"mid"),
("RESULTADO ANTES DO IR",None,"res"),
("(-) Provisão IR (se positivo)",None,"ir"),
("RESULTADO LÍQUIDO",None,"liq"),
("CAIXA ACUMULADO",None,"cx"),
("INVESTIMENTO NECESSÁRIO NO MÊS (aporte)",None,"aporte"),
("CAC (mídia + ligações ÷ clientes novos)",None,"cac"),
]
idx={}
for i,(lab,_,k) in enumerate(lines,2): idx[k]=i;f.cell(i,1,lab)
for j in range(6):
    col=L(2+j);cc=L(4+j);prev=L(1+j)
    def s(k,v): x=f.cell(idx[k],2+j,v);x.number_format=BRL;return x
    x=f.cell(idx["in"],2+j,lines[0][1][j]);x.fill=IN
    f.cell(idx["novos"],2+j,f"={col}{idx['in']}" if j==0 else f"={col}{idx['in']}-{prev}{idx['in']}")
    s("mrr",f"={col}{idx['in']}*{P['ticket']}")
    s("setup",f"={col}{idx['novos']}*{P['setup']}*{P['pctsetup']}")
    s("rec",f"={col}{idx['mrr']}+{col}{idx['setup']}")
    s("var",f"=-{col}{idx['in']}*{P['varcli']}")
    s("gw",f"=-{col}{idx['rec']}*{P['gateway']}")
    s("op",f"=-{CM}{cc}{tot_op}")
    s("mid",f"=-{CM}{cc}{tot_mid}")
    s("res",f"=SUM({col}{idx['rec']}:{col}{idx['mid']})-{col}{idx['rec']}+{col}{idx['rec']}")
    f.cell(idx["res"],2+j).value=f"={col}{idx['rec']}+{col}{idx['var']}+{col}{idx['gw']}+{col}{idx['op']}+{col}{idx['mid']}"
    s("ir",f"=-MAX(0,{col}{idx['res']})*{P['ir']}")
    s("liq",f"={col}{idx['res']}+{col}{idx['ir']}")
    s("cx",f"={col}{idx['liq']}" if j==0 else f"={prev}{idx['cx']}+{col}{idx['liq']}")
    s("aporte",f"=MAX(0,-{col}{idx['liq']})")
    lig_row=conv_end
    s("cac",f"=IF({col}{idx['novos']}>0,(-{col}{idx['mid']}+{CM}{cc}{lig_row})/{col}{idx['novos']},0)")
for k in ["mrr","setup","rec","var","gw","op","mid","res","ir","liq","aporte"]:
    x=f.cell(idx[k],8,f"=SUM(B{idx[k]}:G{idx[k]})");x.number_format=BRL
for k in ["rec","res","liq","cx","aporte"]:
    for j in range(1,9):
        f.cell(idx[k],j).font=B; f.cell(idx[k],j).fill=TOT
f.column_dimensions["A"].width=46
for j in range(2,9): f.column_dimensions[L(j)].width=13
r=idx["cac"]+2
f.cell(r,1,"INVESTIMENTO TOTAL ATÉ SE PAGAR (pior caixa acumulado)").font=Font(bold=True,size=12)
x=f.cell(r,2,f"=-MIN(0,MIN(B{idx['cx']}:G{idx['cx']}))");x.number_format=BRL;x.font=Font(bold=True,size=12);x.fill=PatternFill("solid",fgColor="F8CBAD")
f.cell(r+1,1,"Primeiro mês com resultado líquido positivo")
f.cell(r+1,2,f'=IFERROR(INDEX(B1:G1,MATCH(TRUE,INDEX(B{idx["liq"]}:G{idx["liq"]}>0,0),0)),"após mês 6")')
f.cell(r+3,1,"Células amarelas são editáveis. Clientes ativos seguem o plano do README; se a venda atrasar, reduza aqui e veja o aporte subir.").font=Font(italic=True)
wb.move_sheet("Resultado e caixa",offset=-2)
wb.save("custos-recepcao-ia.xlsx")
